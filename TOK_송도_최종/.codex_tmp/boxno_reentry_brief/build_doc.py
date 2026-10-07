from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(r"D:\Project\TOK_송도_최종\BOX-NO_잔량출고_재입고_처리방안.docx")

INK = "172B4D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5F6B7A"
LIGHT_BLUE = "EAF3F8"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "D9E0E7"
CAUTION = "FFF4CC"
CAUTION_TEXT = "6B5200"
RISK = "9B1C1C"
WHITE = "FFFFFF"
BLACK = "000000"


def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style, size, color=BLACK, bold=False):
    style.font.name = "Arial"
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)


def add_bottom_border(paragraph, color=MID_GRAY, size=8, space=4):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color=MID_GRAY, size=6, inside=True):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    names = ["top", "left", "bottom", "right"]
    if inside:
        names += ["insideH", "insideV"]
    for name in names:
        edge = OxmlElement(f"w:{name}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), str(size))
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), color)
        borders.append(edge)


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def add_real_bullet_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•")
    lvl.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    lvl.append(lvl_jc)

    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.append(ind)
    lvl.append(p_pr)

    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    fonts.set(qn("w:eastAsia"), "맑은 고딕")
    r_pr.append(fonts)
    lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_bullet(doc, num_id, text, bold_prefix=None, after=4):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.167
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    p_pr.insert(0, num_pr)

    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, size=11, bold=True, color=INK)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest, size=11, color=INK)
    else:
        run = p.add_run(text)
        set_run_font(run, size=11, color=INK)
    return p


def add_heading(doc, text, before=10, after=4):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, size=16, bold=True, color=BLUE)
    return p


def add_label_line(doc, label, text, after=3):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    r1 = p.add_run(label)
    set_run_font(r1, size=11, bold=True, color=DARK_BLUE)
    r2 = p.add_run(text)
    set_run_font(r2, size=11, color=INK)
    return p


def add_callout(doc, label, text, fill=LIGHT_BLUE, label_color=DARK_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_table_borders(table, color=fill, size=1, inside=False)
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_margins(cell, top=120, start=160, bottom=120, end=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.10
    r1 = p.add_run(label + "  ")
    set_run_font(r1, size=11, bold=True, color=label_color)
    r2 = p.add_run(text)
    set_run_font(r2, size=11, bold=True, color=INK)
    return table


def set_cell_paragraph(cell, text, *, bold=False, color=INK, align=WD_ALIGN_PARAGRAPH.LEFT, size=9.5):
    p = cell.paragraphs[0]
    clear_paragraph(p)
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)


def build():
    doc = Document()
    section = doc.sections[0]
    # Named Korean-office one-page brief override: A4 while retaining a 6.5 in body width.
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.60)
    section.bottom_margin = Inches(0.58)
    section.left_margin = Inches(0.885)
    section.right_margin = Inches(0.885)
    section.header_distance = Inches(0.30)
    section.footer_distance = Inches(0.30)

    normal = doc.styles["Normal"]
    set_style_font(normal, 11, INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    h1 = doc.styles["Heading 1"]
    set_style_font(h1, 16, BLUE, True)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    set_style_font(h2, 13, BLUE, True)
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(5)
    h2.paragraph_format.keep_with_next = True

    h3 = doc.styles["Heading 3"]
    set_style_font(h3, 12, DARK_BLUE, True)
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True

    bullet_num_id = add_real_bullet_numbering(doc)

    # Quiet running header/footer for a formal one-page decision brief.
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    hr = hp.add_run("업무 협의안 | 2026.09.02")
    set_run_font(hr, size=8.5, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    fr = fp.add_run("검토 근거: 송도 현행 소스 및 제시 업무 시나리오")
    set_run_font(fr, size=8, color=MUTED)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(1)
    tr = title.add_run("잔량 출고 후 재입고 BOX-NO 처리 검토")
    set_run_font(tr, size=23, bold=True, color=BLACK)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(6)
    sr = subtitle.add_run("PLT ID 기반 자동 재입고 전환에 따른 BOX-NO 정합성 검토")
    set_run_font(sr, size=11.5, color=MUTED)
    add_bottom_border(subtitle)

    add_callout(
        doc,
        "핵심 결론",
        "현재 출고 데이터에는 개별 피킹 박스 정보가 없어 자동 계산의 정확성을 보장할 수 없습니다. 업무 규칙이 확정된 BOX-NO 관리 대상만 조건부 자동화해야 합니다.",
    )

    add_heading(doc, "1. 이슈 배경 및 고객 요청", before=8, after=3)
    add_label_line(doc, "AS-IS  ", "PLT ID 없이 잔량 재입고 정보를 작업자가 등록")
    add_label_line(doc, "TO-BE  ", "BCR이 PLT ID를 읽어 MISUBK 잔량으로 자동 재입고")
    add_label_line(
        doc,
        "고객 요청  ",
        "001-012에서 001, 002번 박스를 출고하면 잔량 BOX-NO를 003-012로 자동 변경",
        after=1,
    )

    system_note = doc.add_paragraph()
    system_note.paragraph_format.left_indent = Inches(0.16)
    system_note.paragraph_format.space_before = Pt(2)
    system_note.paragraph_format.space_after = Pt(4)
    system_note.paragraph_format.line_spacing = 1.05
    nr1 = system_note.add_run("소스 확인: ")
    set_run_font(nr1, size=9.5, bold=True, color=RISK)
    nr2 = system_note.add_run(
        "출고 완료 시 MISUBK는 수량만 차감하고 BOX-NO를 유지하며, BCR 재입고도 기존 SUBK_BOXNO를 그대로 사용합니다."
    )
    set_run_font(nr2, size=9.5, color=MUTED)

    add_heading(doc, "2. 현 구조에서 일괄 자동 계산을 보장할 수 없는 이유", before=6, after=3)
    add_bullet(
        doc,
        bullet_num_id,
        "식별 단위 부족: PLT ID는 팔레트 식별자이며, 출고 데이터는 개별 박스가 아닌 수량 중심입니다.",
        bold_prefix="식별 단위 부족: ",
        after=2,
    )
    add_bullet(
        doc,
        bullet_num_id,
        "데이터 의미 차이: OUPT_BOXNO1은 출고 BOX-NO이지 잔량 BOX-NO가 아니므로 MISUBK에 단순 복사할 수 없습니다.",
        bold_prefix="데이터 의미 차이: ",
        after=2,
    )
    add_bullet(
        doc,
        bullet_num_id,
        "범위 표현 한계: 중간 박스가 빠진 불연속 잔량은 단일 BOX-NO 범위로 표현할 수 없습니다.",
        bold_prefix="범위 표현 한계: ",
        after=2,
    )
    add_bullet(
        doc,
        bullet_num_id,
        "환산 기준 부족: BOX당 고정 수량·최저 번호 순차 피킹 규칙이 없으면 수량만으로 잔여 범위를 산출할 수 없습니다.",
        bold_prefix="수량과 박스 수 불일치: ",
        after=2,
    )
    add_bullet(
        doc,
        bullet_num_id,
        "적용 대상 불명확: 비관리 거래처까지 자동 생성하면 실제 사용하지 않는 BOX-NO가 만들어집니다.",
        bold_prefix="적용 대상 불명확: ",
        after=4,
    )

    example = doc.add_paragraph()
    example.paragraph_format.left_indent = Inches(0.16)
    example.paragraph_format.space_before = Pt(1)
    example.paragraph_format.space_after = Pt(3)
    example.paragraph_format.line_spacing = 1.05
    er1 = example.add_run("대표 예외  ")
    set_run_font(er1, size=9.5, bold=True, color=RISK)
    er2 = example.add_run(
        "001, 005 출고 시 잔량은 002-004와 006-012로 분리되어 단일 ‘NNN-NNN’ 범위로 표현할 수 없습니다."
    )
    set_run_font(er2, size=9.5, color=MUTED)

    add_heading(doc, "3. 당사 제안", before=6, after=3)
    add_label_line(
        doc,
        "단기 안정화안  ",
        "실제 출고 완료 시 잔량 BOX-NO를 1회 확인해 MISUBK에 저장합니다. 재입고 시 별도 등록 없이 BCR이 자동 재사용하므로 재입고 작업은 자동화됩니다.",
        after=2,
    )
    add_label_line(
        doc,
        "조건부 자동화안  ",
        "관리 대상·고정 BOX당 수량·최저 번호 순차 피킹·부분/중간 출고 금지 조건을 모두 충족할 때만 계산합니다. 예외는 작업자 확인, 비관리 거래처는 공란을 유지합니다.",
        after=5,
    )

    decision = doc.add_paragraph()
    decision.paragraph_format.left_indent = Inches(0.08)
    decision.paragraph_format.right_indent = Inches(0.08)
    decision.paragraph_format.space_before = Pt(2)
    decision.paragraph_format.space_after = Pt(0)
    decision.paragraph_format.line_spacing = 1.05
    decision_pr = decision._p.get_or_add_pPr()
    decision_shd = OxmlElement("w:shd")
    decision_shd.set(qn("w:fill"), CAUTION)
    decision_pr.append(decision_shd)
    dr1 = decision.add_run("협의 필요  ")
    set_run_font(dr1, size=9.5, bold=True, color=CAUTION_TEXT)
    dr2 = decision.add_run(
        "① 적용 거래처·품목  ② 피킹 순서  ③ BOX당 수량  ④ 중간번호 출고·취소·재포장 예외  ⑤ 비관리 거래처 표기 기준"
    )
    set_run_font(dr2, size=9.5, bold=True, color=INK)

    # Keep all key sections together where practical.
    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True

    doc.core_properties.title = "잔량 출고 후 재입고 BOX-NO 처리 검토"
    doc.core_properties.subject = "PLT ID 기반 자동 재입고 BOX-NO 정합성 검토"
    doc.core_properties.author = "LS E&M"
    doc.core_properties.keywords = "BOX-NO, PLT ID, MISUBK, 잔량 출고, 재입고"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
